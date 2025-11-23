#!/usr/bin/env python3
"""
RoomersAround Rome - Professional Presentation Generator
Creates a professional PowerPoint presentation with mobile app mockups
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from PIL import Image, ImageDraw, ImageFont
import os

# Create images directory
os.makedirs('images', exist_ok=True)

# Color scheme
PRIMARY_BLUE = (66, 135, 245)  # #4287F5
SECONDARY_PURPLE = (155, 89, 182)  # #9B59B6
BACKGROUND_GRAY = (240, 242, 245)  # #F0F2F5
DARK_TEXT = (33, 33, 33)  # #212121
WHITE = (255, 255, 255)

def create_mobile_mockup(screen_name, elements):
    """Create a mobile app mockup image"""
    # Mobile screen dimensions (iPhone-like ratio)
    width, height = 1080, 1920

    # Create image with white background
    img = Image.new('RGB', (width, height), WHITE)
    draw = ImageDraw.Draw(img)

    # Try to use a default font, fallback to basic if not available
    try:
        title_font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 80)
        header_font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 60)
        normal_font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 45)
        small_font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 35)
    except:
        title_font = ImageFont.load_default()
        header_font = ImageFont.load_default()
        normal_font = ImageFont.load_default()
        small_font = ImageFont.load_default()

    # Status bar
    draw.rectangle([(0, 0), (width, 80)], fill=PRIMARY_BLUE)
    draw.text((50, 25), "9:41", fill=WHITE, font=normal_font)

    # Draw elements based on screen type
    if screen_name == "home":
        # Header
        draw.rectangle([(0, 80), (width, 300)], fill=PRIMARY_BLUE)
        draw.text((width//2, 150), "RoomersAround", fill=WHITE, font=title_font, anchor="mm")
        draw.text((width//2, 240), "Find Your Perfect Room in Rome", fill=WHITE, font=small_font, anchor="mm")

        # Search bar
        draw.rounded_rectangle([(100, 350), (width-100, 470)], radius=25, fill=BACKGROUND_GRAY)
        draw.text((150, 410), "🔍 Search location, price...", fill=(150, 150, 150), font=normal_font)

        # Filter buttons
        filters = ["All", "Studio", "1BR", "2BR", "Shared"]
        x_start = 100
        for i, filter_text in enumerate(filters):
            x = x_start + i * 180
            if i == 0:
                draw.rounded_rectangle([(x, 520), (x+150, 600)], radius=15, fill=PRIMARY_BLUE)
                draw.text((x+75, 560), filter_text, fill=WHITE, font=normal_font, anchor="mm")
            else:
                draw.rounded_rectangle([(x, 520), (x+150, 600)], radius=15, outline=PRIMARY_BLUE, width=3)
                draw.text((x+75, 560), filter_text, fill=PRIMARY_BLUE, font=normal_font, anchor="mm")

        # Property cards
        y_pos = 680
        for i in range(3):
            # Card background
            draw.rounded_rectangle([(80, y_pos), (width-80, y_pos+350)], radius=20, fill=WHITE, outline=(220, 220, 220), width=2)

            # Property image placeholder
            draw.rounded_rectangle([(100, y_pos+20), (width-100, y_pos+220)], radius=15, fill=BACKGROUND_GRAY)
            draw.text((width//2, y_pos+120), "🏠", fill=PRIMARY_BLUE, font=title_font, anchor="mm")

            # Property details
            draw.text((110, y_pos+250), f"Cozy Apartment in Trastevere", fill=DARK_TEXT, font=header_font)
            draw.text((110, y_pos+310), f"€850/month • 45m²", fill=(100, 100, 100), font=normal_font)

            y_pos += 380

    elif screen_name == "listings":
        # Header
        draw.rectangle([(0, 80), (width, 250)], fill=PRIMARY_BLUE)
        draw.text((50, 150), "← Back", fill=WHITE, font=normal_font)
        draw.text((width//2, 165), "Available Properties", fill=WHITE, font=header_font, anchor="mm")

        # List items
        y_pos = 280
        for i in range(5):
            # List item background
            draw.rectangle([(0, y_pos), (width, y_pos+280)], fill=WHITE)
            draw.line([(0, y_pos+280), (width, y_pos+280)], fill=(230, 230, 230), width=2)

            # Thumbnail
            draw.rounded_rectangle([(40, y_pos+20), (280, y_pos+260)], radius=15, fill=BACKGROUND_GRAY)
            draw.text((160, y_pos+140), "🏠", fill=PRIMARY_BLUE, font=header_font, anchor="mm")

            # Details
            draw.text((320, y_pos+40), f"Modern Studio", fill=DARK_TEXT, font=header_font)
            draw.text((320, y_pos+100), f"Testaccio District", fill=(120, 120, 120), font=normal_font)
            draw.text((320, y_pos+150), f"€750/month", fill=PRIMARY_BLUE, font=header_font)
            draw.text((320, y_pos+200), "✓ Verified • ★ 4.8", fill=(80, 200, 120), font=normal_font)

            y_pos += 280

    elif screen_name == "details":
        # Image gallery
        draw.rectangle([(0, 80), (width, 700)], fill=BACKGROUND_GRAY)
        draw.text((width//2, 390), "🏠 Property Gallery", fill=PRIMARY_BLUE, font=title_font, anchor="mm")

        # Back button
        draw.rounded_rectangle([(40, 120), (160, 200)], radius=20, fill=(0, 0, 0, 128))
        draw.text((100, 160), "←", fill=WHITE, font=header_font, anchor="mm")

        # Favorite button
        draw.rounded_rectangle([(width-160, 120), (width-40, 200)], radius=20, fill=(0, 0, 0, 128))
        draw.text((width-100, 160), "♥", fill=WHITE, font=header_font, anchor="mm")

        # Property info
        draw.text((60, 750), "Beautiful Apartment in Monti", fill=DARK_TEXT, font=title_font)
        draw.text((60, 850), "Via dei Serpenti, Monti, Rome", fill=(120, 120, 120), font=normal_font)

        # Price tag
        draw.rounded_rectangle([(60, 920), (350, 1010)], radius=15, fill=PRIMARY_BLUE)
        draw.text((205, 965), "€950/month", fill=WHITE, font=header_font, anchor="mm")

        # Features
        features = [
            ("📏", "65 m²"),
            ("🛏️", "2 Bedrooms"),
            ("🚿", "1 Bathroom"),
            ("📍", "Floor 3")
        ]

        x_start = 60
        y_start = 1080
        for i, (icon, text) in enumerate(features):
            x = x_start + (i % 2) * 500
            y = y_start + (i // 2) * 120
            draw.text((x, y), f"{icon} {text}", fill=DARK_TEXT, font=normal_font)

        # Description
        draw.text((60, 1350), "Description", fill=DARK_TEXT, font=header_font)
        draw.text((60, 1430), "Charming apartment in the heart of Monti,", fill=(80, 80, 80), font=normal_font)
        draw.text((60, 1490), "close to Colosseum. Fully furnished with", fill=(80, 80, 80), font=normal_font)
        draw.text((60, 1550), "modern amenities and great transport links.", fill=(80, 80, 80), font=normal_font)

        # Action buttons
        draw.rounded_rectangle([(60, 1700), (width-60, 1820)], radius=20, fill=PRIMARY_BLUE)
        draw.text((width//2, 1760), "Contact Owner", fill=WHITE, font=header_font, anchor="mm")

    elif screen_name == "profile":
        # Header
        draw.rectangle([(0, 80), (width, 400)], fill=PRIMARY_BLUE)

        # Profile picture
        draw.ellipse([(width//2-150, 200), (width//2+150, 500)], fill=WHITE)
        draw.text((width//2, 350), "👤", fill=PRIMARY_BLUE, font=title_font, anchor="mm")

        # User name
        draw.text((width//2, 550), "Dolly Mathur", fill=DARK_TEXT, font=title_font, anchor="mm")
        draw.text((width//2, 640), "Member since 2025", fill=(120, 120, 120), font=normal_font, anchor="mm")

        # Stats
        stats = [("12", "Listings"), ("45", "Reviews"), ("4.9", "Rating")]
        x_start = 100
        for i, (number, label) in enumerate(stats):
            x = x_start + i * 300
            draw.text((x+100, 750), number, fill=PRIMARY_BLUE, font=title_font, anchor="mm")
            draw.text((x+100, 840), label, fill=(80, 80, 80), font=normal_font, anchor="mm")

        # Menu items
        menu_items = [
            ("🏠", "My Listings"),
            ("💬", "Messages"),
            ("⭐", "Favorites"),
            ("⚙️", "Settings"),
            ("❓", "Help & Support")
        ]

        y_pos = 950
        for icon, text in menu_items:
            draw.rectangle([(60, y_pos), (width-60, y_pos+120)], fill=WHITE, outline=(230, 230, 230), width=2)
            draw.text((100, y_pos+60), f"{icon}  {text}", fill=DARK_TEXT, font=header_font)
            draw.text((width-100, y_pos+60), "›", fill=(180, 180, 180), font=header_font, anchor="mm")
            y_pos += 140

    elif screen_name == "booking":
        # Header
        draw.rectangle([(0, 80), (width, 250)], fill=PRIMARY_BLUE)
        draw.text((50, 150), "← Back", fill=WHITE, font=normal_font)
        draw.text((width//2, 165), "Book Your Stay", fill=WHITE, font=header_font, anchor="mm")

        # Property summary
        draw.rounded_rectangle([(60, 300), (width-60, 550)], radius=20, fill=BACKGROUND_GRAY)
        draw.text((100, 350), "Apartment in Monti", fill=DARK_TEXT, font=header_font)
        draw.text((100, 420), "€950/month", fill=PRIMARY_BLUE, font=header_font)
        draw.text((100, 490), "Available from: 01 Jan 2025", fill=(100, 100, 100), font=normal_font)

        # Form fields
        y_pos = 620
        fields = ["Move-in Date", "Duration (months)", "Number of Occupants"]
        for field in fields:
            draw.text((80, y_pos), field, fill=(80, 80, 80), font=normal_font)
            draw.rounded_rectangle([(80, y_pos+50), (width-80, y_pos+150)], radius=15, fill=WHITE, outline=(200, 200, 200), width=2)
            y_pos += 220

        # Cost summary
        draw.text((80, 1350), "Summary", fill=DARK_TEXT, font=header_font)
        draw.rectangle([(80, 1420), (width-80, 1650)], fill=BACKGROUND_GRAY)
        draw.text((100, 1460), "Monthly Rent", fill=(80, 80, 80), font=normal_font)
        draw.text((width-100, 1460), "€950", fill=DARK_TEXT, font=normal_font, anchor="rm")
        draw.text((100, 1530), "Service Fee", fill=(80, 80, 80), font=normal_font)
        draw.text((width-100, 1530), "€50", fill=DARK_TEXT, font=normal_font, anchor="rm")
        draw.line([(100, 1580), (width-100, 1580)], fill=(180, 180, 180), width=2)
        draw.text((100, 1600), "Total", fill=DARK_TEXT, font=header_font)
        draw.text((width-100, 1600), "€1000", fill=PRIMARY_BLUE, font=header_font, anchor="rm")

        # Submit button
        draw.rounded_rectangle([(80, 1730), (width-80, 1850)], radius=20, fill=PRIMARY_BLUE)
        draw.text((width//2, 1790), "Submit Request", fill=WHITE, font=header_font, anchor="mm")

    # Save image
    img.save(f'images/mockup_{screen_name}.png')
    print(f"Created mockup: mockup_{screen_name}.png")

def create_presentation():
    """Create the professional PowerPoint presentation"""
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # Slide 1: Title Slide
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # Background
    background = slide.shapes.add_shape(
        1,  # Rectangle
        0, 0, prs.slide_width, prs.slide_height
    )
    background.fill.solid()
    background.fill.fore_color.rgb = RGBColor(240, 242, 245)
    background.line.fill.background()

    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2), Inches(9), Inches(1.5))
    title_frame = title_box.text_frame
    title_frame.text = "RoomersAround Rome"
    title_para = title_frame.paragraphs[0]
    title_para.font.size = Pt(60)
    title_para.font.bold = True
    title_para.font.color.rgb = RGBColor(*PRIMARY_BLUE)
    title_para.alignment = PP_ALIGN.CENTER

    # Subtitle
    subtitle_box = slide.shapes.add_textbox(Inches(0.5), Inches(3.5), Inches(9), Inches(1))
    subtitle_frame = subtitle_box.text_frame
    subtitle_frame.text = "Property Rental Platform with Community Focus"
    subtitle_para = subtitle_frame.paragraphs[0]
    subtitle_para.font.size = Pt(32)
    subtitle_para.font.color.rgb = RGBColor(*SECONDARY_PURPLE)
    subtitle_para.alignment = PP_ALIGN.CENTER

    # Tagline
    tagline_box = slide.shapes.add_textbox(Inches(0.5), Inches(4.8), Inches(9), Inches(0.8))
    tagline_frame = tagline_box.text_frame
    tagline_frame.text = "Connecting Property Seekers & Owners in the Eternal City"
    tagline_para = tagline_frame.paragraphs[0]
    tagline_para.font.size = Pt(20)
    tagline_para.font.color.rgb = RGBColor(*DARK_TEXT)
    tagline_para.alignment = PP_ALIGN.CENTER

    # Presenter
    presenter_box = slide.shapes.add_textbox(Inches(0.5), Inches(6.5), Inches(9), Inches(0.6))
    presenter_frame = presenter_box.text_frame
    presenter_frame.text = "Presented by: Dolly Mathur | 2025"
    presenter_para = presenter_frame.paragraphs[0]
    presenter_para.font.size = Pt(18)
    presenter_para.font.color.rgb = RGBColor(80, 80, 80)
    presenter_para.alignment = PP_ALIGN.CENTER

    # Slide 2: Problem Statement
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    background.fill.solid()
    background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    background.line.fill.background()

    # Header
    header_shape = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(1))
    header_shape.fill.solid()
    header_shape.fill.fore_color.rgb = RGBColor(*PRIMARY_BLUE)
    header_shape.line.fill.background()

    header_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(9), Inches(0.5))
    header_frame = header_box.text_frame
    header_frame.text = "The Problem"
    header_para = header_frame.paragraphs[0]
    header_para.font.size = Pt(36)
    header_para.font.bold = True
    header_para.font.color.rgb = RGBColor(255, 255, 255)

    # Content
    content_box = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(8), Inches(5))
    content_frame = content_box.text_frame
    content_frame.word_wrap = True

    problems = [
        "Finding quality rental properties in Rome is challenging and time-consuming",
        "Lack of trust between property seekers and owners",
        "Limited community engagement in existing rental platforms",
        "Difficulty in verifying property authenticity and owner credibility",
        "Language barriers for international students and expats"
    ]

    for problem in problems:
        p = content_frame.add_paragraph()
        p.text = problem
        p.font.size = Pt(24)
        p.font.color.rgb = RGBColor(*DARK_TEXT)
        p.level = 0
        p.space_before = Pt(18)

    # Slide 3: Our Solution
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    background.fill.solid()
    background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    background.line.fill.background()

    header_shape = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(1))
    header_shape.fill.solid()
    header_shape.fill.fore_color.rgb = RGBColor(*PRIMARY_BLUE)
    header_shape.line.fill.background()

    header_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(9), Inches(0.5))
    header_frame = header_box.text_frame
    header_frame.text = "Our Solution"
    header_para = header_frame.paragraphs[0]
    header_para.font.size = Pt(36)
    header_para.font.bold = True
    header_para.font.color.rgb = RGBColor(255, 255, 255)

    # Main description
    desc_box = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(8), Inches(1.5))
    desc_frame = desc_box.text_frame
    desc_frame.word_wrap = True
    desc_frame.text = "RoomersAround Rome is a comprehensive property rental platform that creates a trusted community ecosystem connecting property seekers with verified owners in Rome."
    desc_para = desc_frame.paragraphs[0]
    desc_para.font.size = Pt(22)
    desc_para.font.color.rgb = RGBColor(*DARK_TEXT)
    desc_para.alignment = PP_ALIGN.CENTER

    # Key differentiators
    diff_box = slide.shapes.add_textbox(Inches(1), Inches(3.3), Inches(8), Inches(3.5))
    diff_frame = diff_box.text_frame
    diff_frame.word_wrap = True

    differentiators = [
        ("Verified Listings:", "All properties undergo authentication checks"),
        ("Community Reviews:", "Transparent rating system for properties and users"),
        ("Cultural Integration:", "Multilingual support and local area guides"),
        ("Smart Matching:", "AI-powered recommendations based on preferences")
    ]

    for title, desc in differentiators:
        p = diff_frame.add_paragraph()
        p.text = title
        p.font.size = Pt(20)
        p.font.bold = True
        p.font.color.rgb = RGBColor(*PRIMARY_BLUE)
        p.space_before = Pt(15)

        p2 = diff_frame.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(18)
        p2.font.color.rgb = RGBColor(*DARK_TEXT)
        p2.level = 1

    # Slide 4: Key Features
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    background.fill.solid()
    background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    background.line.fill.background()

    header_shape = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(1))
    header_shape.fill.solid()
    header_shape.fill.fore_color.rgb = RGBColor(*PRIMARY_BLUE)
    header_shape.line.fill.background()

    header_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(9), Inches(0.5))
    header_frame = header_box.text_frame
    header_frame.text = "Key Features"
    header_para = header_frame.paragraphs[0]
    header_para.font.size = Pt(36)
    header_para.font.bold = True
    header_para.font.color.rgb = RGBColor(255, 255, 255)

    # Features in two columns
    features_left = [
        ("Advanced Search & Filters", "Location, price, amenities, availability"),
        ("Virtual Property Tours", "360° views and video walkthroughs"),
        ("Secure Messaging", "In-app communication between parties"),
        ("Document Management", "Digital contracts and paperwork")
    ]

    features_right = [
        ("Payment Integration", "Secure online rent payments"),
        ("Community Forum", "Connect with fellow renters"),
        ("Local Insights", "Neighborhood guides and tips"),
        ("Favorites & Alerts", "Save properties and get notifications")
    ]

    # Left column
    y_pos = 1.3
    for title, desc in features_left:
        # Feature box
        box = slide.shapes.add_shape(1, Inches(0.3), Inches(y_pos), Inches(4.5), Inches(1.2))
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(*BACKGROUND_GRAY)
        box.line.color.rgb = RGBColor(*PRIMARY_BLUE)
        box.line.width = Pt(2)

        text_box = slide.shapes.add_textbox(Inches(0.5), Inches(y_pos + 0.15), Inches(4), Inches(0.9))
        text_frame = text_box.text_frame
        text_frame.word_wrap = True

        p1 = text_frame.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(18)
        p1.font.bold = True
        p1.font.color.rgb = RGBColor(*PRIMARY_BLUE)

        p2 = text_frame.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(14)
        p2.font.color.rgb = RGBColor(*DARK_TEXT)

        y_pos += 1.4

    # Right column
    y_pos = 1.3
    for title, desc in features_right:
        box = slide.shapes.add_shape(1, Inches(5.2), Inches(y_pos), Inches(4.5), Inches(1.2))
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(*BACKGROUND_GRAY)
        box.line.color.rgb = RGBColor(*PRIMARY_BLUE)
        box.line.width = Pt(2)

        text_box = slide.shapes.add_textbox(Inches(5.4), Inches(y_pos + 0.15), Inches(4), Inches(0.9))
        text_frame = text_box.text_frame
        text_frame.word_wrap = True

        p1 = text_frame.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(18)
        p1.font.bold = True
        p1.font.color.rgb = RGBColor(*PRIMARY_BLUE)

        p2 = text_frame.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(14)
        p2.font.color.rgb = RGBColor(*DARK_TEXT)

        y_pos += 1.4

    # Slide 5: Technology Stack
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    background.fill.solid()
    background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    background.line.fill.background()

    header_shape = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(1))
    header_shape.fill.solid()
    header_shape.fill.fore_color.rgb = RGBColor(*PRIMARY_BLUE)
    header_shape.line.fill.background()

    header_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(9), Inches(0.5))
    header_frame = header_box.text_frame
    header_frame.text = "Technology Stack"
    header_para = header_frame.paragraphs[0]
    header_para.font.size = Pt(36)
    header_para.font.bold = True
    header_para.font.color.rgb = RGBColor(255, 255, 255)

    # Tech categories
    tech_stack = [
        ("Frontend", ["React Native", "Redux for state management", "Tailwind CSS for styling"]),
        ("Backend", ["Node.js with Express.js", "RESTful API architecture", "JWT authentication"]),
        ("Database", ["MongoDB for data storage", "Redis for caching", "AWS S3 for media files"]),
        ("Mobile", ["iOS & Android native support", "Push notifications", "Geolocation services"]),
        ("Infrastructure", ["AWS Cloud hosting", "CI/CD with GitHub Actions", "Docker containerization"])
    ]

    y_pos = 1.3
    for category, technologies in tech_stack:
        # Category title
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(y_pos), Inches(3), Inches(0.4))
        cat_frame = cat_box.text_frame
        cat_frame.text = category
        cat_para = cat_frame.paragraphs[0]
        cat_para.font.size = Pt(22)
        cat_para.font.bold = True
        cat_para.font.color.rgb = RGBColor(*PRIMARY_BLUE)

        # Technologies
        tech_box = slide.shapes.add_textbox(Inches(3.5), Inches(y_pos), Inches(6), Inches(0.8))
        tech_frame = tech_box.text_frame
        tech_frame.word_wrap = True

        for tech in technologies:
            p = tech_frame.add_paragraph() if tech != technologies[0] else tech_frame.paragraphs[0]
            p.text = f"• {tech}"
            p.font.size = Pt(16)
            p.font.color.rgb = RGBColor(*DARK_TEXT)

        y_pos += 1.1

    # Slide 6: Mobile App - Home Screen
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    background.fill.solid()
    background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    background.line.fill.background()

    header_shape = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(1))
    header_shape.fill.solid()
    header_shape.fill.fore_color.rgb = RGBColor(*PRIMARY_BLUE)
    header_shape.line.fill.background()

    header_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(9), Inches(0.5))
    header_frame = header_box.text_frame
    header_frame.text = "Mobile App - Home Screen"
    header_para = header_frame.paragraphs[0]
    header_para.font.size = Pt(36)
    header_para.font.bold = True
    header_para.font.color.rgb = RGBColor(255, 255, 255)

    # Add mockup image
    if os.path.exists('images/mockup_home.png'):
        slide.shapes.add_picture('images/mockup_home.png', Inches(3), Inches(1.3), height=Inches(6))

    # Description
    desc_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(2.3), Inches(5.5))
    desc_frame = desc_box.text_frame
    desc_frame.word_wrap = True

    desc_frame.text = "Features:\n\n"

    features = [
        "Smart search bar",
        "Quick filters",
        "Property cards with key info",
        "Image previews",
        "Instant access to listings"
    ]

    for feature in features:
        p = desc_frame.add_paragraph()
        p.text = f"• {feature}"
        p.font.size = Pt(18)
        p.font.color.rgb = RGBColor(*DARK_TEXT)
        p.space_before = Pt(8)

    # Slide 7: Mobile App - Property Listings
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    background.fill.solid()
    background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    background.line.fill.background()

    header_shape = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(1))
    header_shape.fill.solid()
    header_shape.fill.fore_color.rgb = RGBColor(*PRIMARY_BLUE)
    header_shape.line.fill.background()

    header_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(9), Inches(0.5))
    header_frame = header_box.text_frame
    header_frame.text = "Mobile App - Property Listings"
    header_para = header_frame.paragraphs[0]
    header_para.font.size = Pt(36)
    header_para.font.bold = True
    header_para.font.color.rgb = RGBColor(255, 255, 255)

    if os.path.exists('images/mockup_listings.png'):
        slide.shapes.add_picture('images/mockup_listings.png', Inches(3), Inches(1.3), height=Inches(6))

    desc_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(2.3), Inches(5.5))
    desc_frame = desc_box.text_frame
    desc_frame.word_wrap = True
    desc_frame.text = "Features:\n\n"

    features = [
        "List view of properties",
        "Thumbnail previews",
        "Key property details",
        "Pricing information",
        "Verification badges",
        "Rating indicators"
    ]

    for feature in features:
        p = desc_frame.add_paragraph()
        p.text = f"• {feature}"
        p.font.size = Pt(18)
        p.font.color.rgb = RGBColor(*DARK_TEXT)
        p.space_before = Pt(8)

    # Slide 8: Mobile App - Property Details
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    background.fill.solid()
    background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    background.line.fill.background()

    header_shape = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(1))
    header_shape.fill.solid()
    header_shape.fill.fore_color.rgb = RGBColor(*PRIMARY_BLUE)
    header_shape.line.fill.background()

    header_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(9), Inches(0.5))
    header_frame = header_box.text_frame
    header_frame.text = "Mobile App - Property Details"
    header_para = header_frame.paragraphs[0]
    header_para.font.size = Pt(36)
    header_para.font.bold = True
    header_para.font.color.rgb = RGBColor(255, 255, 255)

    if os.path.exists('images/mockup_details.png'):
        slide.shapes.add_picture('images/mockup_details.png', Inches(3), Inches(1.3), height=Inches(6))

    desc_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(2.3), Inches(5.5))
    desc_frame = desc_box.text_frame
    desc_frame.word_wrap = True
    desc_frame.text = "Features:\n\n"

    features = [
        "Photo gallery",
        "Full property details",
        "Room specifications",
        "Amenities list",
        "Neighborhood info",
        "Contact owner button",
        "Save to favorites"
    ]

    for feature in features:
        p = desc_frame.add_paragraph()
        p.text = f"• {feature}"
        p.font.size = Pt(18)
        p.font.color.rgb = RGBColor(*DARK_TEXT)
        p.space_before = Pt(6)

    # Slide 9: Mobile App - User Profile
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    background.fill.solid()
    background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    background.line.fill.background()

    header_shape = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(1))
    header_shape.fill.solid()
    header_shape.fill.fore_color.rgb = RGBColor(*PRIMARY_BLUE)
    header_shape.line.fill.background()

    header_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(9), Inches(0.5))
    header_frame = header_box.text_frame
    header_frame.text = "Mobile App - User Profile"
    header_para = header_frame.paragraphs[0]
    header_para.font.size = Pt(36)
    header_para.font.bold = True
    header_para.font.color.rgb = RGBColor(255, 255, 255)

    if os.path.exists('images/mockup_profile.png'):
        slide.shapes.add_picture('images/mockup_profile.png', Inches(3), Inches(1.3), height=Inches(6))

    desc_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(2.3), Inches(5.5))
    desc_frame = desc_box.text_frame
    desc_frame.word_wrap = True
    desc_frame.text = "Features:\n\n"

    features = [
        "User profile",
        "Activity statistics",
        "My listings",
        "Messages inbox",
        "Saved favorites",
        "Account settings",
        "Help & support"
    ]

    for feature in features:
        p = desc_frame.add_paragraph()
        p.text = f"• {feature}"
        p.font.size = Pt(18)
        p.font.color.rgb = RGBColor(*DARK_TEXT)
        p.space_before = Pt(6)

    # Slide 10: Mobile App - Booking
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    background.fill.solid()
    background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    background.line.fill.background()

    header_shape = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(1))
    header_shape.fill.solid()
    header_shape.fill.fore_color.rgb = RGBColor(*PRIMARY_BLUE)
    header_shape.line.fill.background()

    header_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(9), Inches(0.5))
    header_frame = header_box.text_frame
    header_frame.text = "Mobile App - Booking Flow"
    header_para = header_frame.paragraphs[0]
    header_para.font.size = Pt(36)
    header_para.font.bold = True
    header_para.font.color.rgb = RGBColor(255, 255, 255)

    if os.path.exists('images/mockup_booking.png'):
        slide.shapes.add_picture('images/mockup_booking.png', Inches(3), Inches(1.3), height=Inches(6))

    desc_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(2.3), Inches(5.5))
    desc_frame = desc_box.text_frame
    desc_frame.word_wrap = True
    desc_frame.text = "Features:\n\n"

    features = [
        "Booking request form",
        "Date selection",
        "Duration picker",
        "Cost breakdown",
        "Service fee display",
        "Total calculation",
        "Submit request"
    ]

    for feature in features:
        p = desc_frame.add_paragraph()
        p.text = f"• {feature}"
        p.font.size = Pt(18)
        p.font.color.rgb = RGBColor(*DARK_TEXT)
        p.space_before = Pt(6)

    # Slide 11: Future Roadmap
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    background.fill.solid()
    background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    background.line.fill.background()

    header_shape = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(1))
    header_shape.fill.solid()
    header_shape.fill.fore_color.rgb = RGBColor(*PRIMARY_BLUE)
    header_shape.line.fill.background()

    header_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(9), Inches(0.5))
    header_frame = header_box.text_frame
    header_frame.text = "Future Roadmap"
    header_para = header_frame.paragraphs[0]
    header_para.font.size = Pt(36)
    header_para.font.bold = True
    header_para.font.color.rgb = RGBColor(255, 255, 255)

    # Roadmap phases
    phases = [
        ("Phase 1: Q1 2025", [
            "Launch MVP for Rome",
            "Basic search and listing features",
            "User authentication and profiles",
            "Property verification system"
        ]),
        ("Phase 2: Q2 2025", [
            "AI-powered recommendations",
            "Virtual property tours",
            "Integrated payment system",
            "Mobile app optimization"
        ]),
        ("Phase 3: Q3 2025", [
            "Community forum and events",
            "Multi-language support (Italian, English, Spanish)",
            "Advanced analytics dashboard",
            "Expand to other Italian cities"
        ]),
        ("Phase 4: Q4 2025", [
            "AR property viewing",
            "Blockchain-based contracts",
            "Partner integrations (utilities, moving services)",
            "European expansion"
        ])
    ]

    y_pos = 1.3
    for phase, features in phases:
        # Phase box
        phase_box = slide.shapes.add_shape(1, Inches(0.5), Inches(y_pos), Inches(9), Inches(1.2))
        phase_box.fill.solid()
        phase_box.fill.fore_color.rgb = RGBColor(*BACKGROUND_GRAY)
        phase_box.line.color.rgb = RGBColor(*SECONDARY_PURPLE)
        phase_box.line.width = Pt(3)

        # Phase title
        title_box = slide.shapes.add_textbox(Inches(0.7), Inches(y_pos + 0.1), Inches(8.5), Inches(0.3))
        title_frame = title_box.text_frame
        title_frame.text = phase
        title_para = title_frame.paragraphs[0]
        title_para.font.size = Pt(20)
        title_para.font.bold = True
        title_para.font.color.rgb = RGBColor(*SECONDARY_PURPLE)

        # Features
        feat_box = slide.shapes.add_textbox(Inches(0.7), Inches(y_pos + 0.45), Inches(8.5), Inches(0.65))
        feat_frame = feat_box.text_frame
        feat_frame.word_wrap = True
        feat_frame.text = " • ".join(features)
        feat_para = feat_frame.paragraphs[0]
        feat_para.font.size = Pt(14)
        feat_para.font.color.rgb = RGBColor(*DARK_TEXT)

        y_pos += 1.4

    # Slide 12: Market Opportunity
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    background.fill.solid()
    background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    background.line.fill.background()

    header_shape = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(1))
    header_shape.fill.solid()
    header_shape.fill.fore_color.rgb = RGBColor(*PRIMARY_BLUE)
    header_shape.line.fill.background()

    header_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(9), Inches(0.5))
    header_frame = header_box.text_frame
    header_frame.text = "Market Opportunity"
    header_para = header_frame.paragraphs[0]
    header_para.font.size = Pt(36)
    header_para.font.bold = True
    header_para.font.color.rgb = RGBColor(255, 255, 255)

    # Stats boxes
    stats = [
        ("50,000+", "Students in Rome", SECONDARY_PURPLE),
        ("€15B", "Italian rental market", PRIMARY_BLUE),
        ("30%", "Growth in expat population", SECONDARY_PURPLE),
        ("2.8M", "Rome population", PRIMARY_BLUE)
    ]

    x_positions = [0.5, 5.2, 0.5, 5.2]
    y_positions = [1.5, 1.5, 3.5, 3.5]

    for i, (stat, label, color) in enumerate(stats):
        stat_box = slide.shapes.add_shape(1, Inches(x_positions[i]), Inches(y_positions[i]), Inches(4.3), Inches(1.5))
        stat_box.fill.solid()
        stat_box.fill.fore_color.rgb = RGBColor(*color)
        stat_box.line.fill.background()

        # Stat number
        num_box = slide.shapes.add_textbox(Inches(x_positions[i]), Inches(y_positions[i] + 0.25), Inches(4.3), Inches(0.6))
        num_frame = num_box.text_frame
        num_frame.text = stat
        num_para = num_frame.paragraphs[0]
        num_para.font.size = Pt(48)
        num_para.font.bold = True
        num_para.font.color.rgb = RGBColor(255, 255, 255)
        num_para.alignment = PP_ALIGN.CENTER

        # Label
        label_box = slide.shapes.add_textbox(Inches(x_positions[i]), Inches(y_positions[i] + 0.9), Inches(4.3), Inches(0.4))
        label_frame = label_box.text_frame
        label_frame.text = label
        label_para = label_frame.paragraphs[0]
        label_para.font.size = Pt(20)
        label_para.font.color.rgb = RGBColor(255, 255, 255)
        label_para.alignment = PP_ALIGN.CENTER

    # Key insights
    insight_box = slide.shapes.add_textbox(Inches(0.8), Inches(5.5), Inches(8.4), Inches(1.5))
    insight_frame = insight_box.text_frame
    insight_frame.word_wrap = True

    p = insight_frame.paragraphs[0]
    p.text = "Key Insight:"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = RGBColor(*PRIMARY_BLUE)

    p2 = insight_frame.add_paragraph()
    p2.text = "Rome attracts thousands of students and young professionals annually, creating a consistent demand for quality rental housing. RoomersAround Rome bridges the gap between supply and demand with a trusted, community-driven platform."
    p2.font.size = Pt(18)
    p2.font.color.rgb = RGBColor(*DARK_TEXT)
    p2.space_before = Pt(10)

    # Slide 13: Closing / Thank You
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    background.fill.solid()
    background.fill.fore_color.rgb = RGBColor(*PRIMARY_BLUE)
    background.line.fill.background()

    # Thank you text
    thank_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(9), Inches(1))
    thank_frame = thank_box.text_frame
    thank_frame.text = "Thank You!"
    thank_para = thank_frame.paragraphs[0]
    thank_para.font.size = Pt(72)
    thank_para.font.bold = True
    thank_para.font.color.rgb = RGBColor(255, 255, 255)
    thank_para.alignment = PP_ALIGN.CENTER

    # Tagline
    tagline_box = slide.shapes.add_textbox(Inches(0.5), Inches(3.8), Inches(9), Inches(0.8))
    tagline_frame = tagline_box.text_frame
    tagline_frame.text = "Let's Build the Future of Property Rental Together"
    tagline_para = tagline_frame.paragraphs[0]
    tagline_para.font.size = Pt(28)
    tagline_para.font.color.rgb = RGBColor(255, 255, 255)
    tagline_para.alignment = PP_ALIGN.CENTER

    # Contact info
    contact_box = slide.shapes.add_textbox(Inches(0.5), Inches(5.5), Inches(9), Inches(1))
    contact_frame = contact_box.text_frame

    p1 = contact_frame.paragraphs[0]
    p1.text = "RoomersAround Rome"
    p1.font.size = Pt(24)
    p1.font.color.rgb = RGBColor(255, 255, 255)
    p1.alignment = PP_ALIGN.CENTER

    p2 = contact_frame.add_paragraph()
    p2.text = "Presented by: Dolly Mathur"
    p2.font.size = Pt(20)
    p2.font.color.rgb = RGBColor(255, 255, 255)
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(10)

    p3 = contact_frame.add_paragraph()
    p3.text = "2025"
    p3.font.size = Pt(20)
    p3.font.color.rgb = RGBColor(255, 255, 255)
    p3.alignment = PP_ALIGN.CENTER
    p3.space_before = Pt(5)

    # Save presentation
    prs.save('RoomersAround_Rome_Presentation.pptx')
    print("Presentation created successfully: RoomersAround_Rome_Presentation.pptx")

if __name__ == "__main__":
    print("Creating mobile app mockups...")
    create_mobile_mockup("home", {})
    create_mobile_mockup("listings", {})
    create_mobile_mockup("details", {})
    create_mobile_mockup("profile", {})
    create_mobile_mockup("booking", {})

    print("\nCreating PowerPoint presentation...")
    create_presentation()

    print("\n✓ All done! Your professional presentation is ready.")
    print("  - Presentation file: RoomersAround_Rome_Presentation.pptx")
    print("  - Mockup images: images/mockup_*.png")
